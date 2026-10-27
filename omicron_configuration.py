# -*- coding: utf-8 -*-
"""
omicron_configuration.py - OMICRON Test Universe hardware configurations (.ohc) from the
General_PB project (author: German Ievgen).

  python omicron_configuration.py                      -> asks for the project folder
  python omicron_configuration.py "D:\\Projects\\Bucha"  -> project folder given
  options: --device W1G_F60 (repeatable) | --all | --master default_01.ohc

Project folder:
  General_PB.xlsm         sheet SV  (SV ID -> IEDName, MAC, APP ID, VLAN, VLAN p, Dataset)
  Library\\*.ohc           master file(s): GOOSE 7x24 virtual I/O + 3 SV streams
  DOC\\<IED>.xlsm          sheets Omicron Binary Inputs / Outputs, Omicron AI <test>
Result (folder Omicron next to DOC, one sub-folder per IED):
  Omicron\\<IED>\\<IED>_Omicron_AI_Test_<n>.ohc   (one per test sheet with SV data)
  Omicron\\<IED>\\<IED>_Omicron.ohc               (device without SV data: GOOSE only)
  Omicron\\_omicron_report.txt
Also used by the GUI ied2omicron.py (functions list_devices / run).

What is written into the master (everything else stays as it is):
  binary inputs  : name of "Bin. in n" / "vBin. in g-k"   <- Omicron Binary Inputs, column D (Omicron ID)
  binary outputs : name of "Bin. out n" / "vBin. out g-k" <- Omicron Binary Outputs, column D
  analog outputs : SV Voltage 1..3 / SV Current 1..3 (4 channels each) <- Omicron AI <test>, column D
                   V L1..V E -> Voltage 1, V(2)-x -> Voltage 2, V(3)-x -> Voltage 3 (I the same)
  name flag      : every renamed channel gets the flag 'name changed by the user' (as Test Universe)
  SV stream n    : SV ID of group n (column F of V(n) / I(n)) + General_PB SV, written into the SV
                   configuration of the stream (dialog 'Configure Sampled Values', tab Stream n):
                   svID, MAC, APP ID, VLAN, VLAN p, ConfRev 1, and the Dataset Mapping (8 dataset
                   attributes -> CMC channels / names); a stream without SV ID stays as in the master
                   v0.0.4 - as done by hand in Test Universe, these stay AS IN THE MASTER:
                   Control block reference, Dataset reference, the tables 'Merging Unit' and
                   'Data Attributes' (report); time source / PTP come from the master too
An empty Omicron ID keeps the name of the master ("vBin. in 1-2", "V(2)-1", ...).
Needs: Python 3, openpyxl (pip install openpyxl).
"""
import os, re, sys, glob, shutil, struct, tempfile, datetime, argparse, warnings

warnings.filterwarnings("ignore")
try:
    import openpyxl
except ImportError:
    print("openpyxl is missing: pip install openpyxl")
    sys.exit(1)

BI_SHEET = "Omicron Binary Inputs"
BO_SHEET = "Omicron Binary Outputs"
AI_PREFIX = "Omicron AI "
OUT_DIR = "Omicron"                 # next to DOC; one sub-folder per IED
REPORT = "_omicron_report.txt"
ETH_PORT = "ETH1"
CONF_REV = 1

MU_FIELDS = ["Index", "Enabled", "Ied", "SvId", "Simulated", "EthernetPort", "SclPath", "ImportTime",
             "DestinationMacAddress", "ControlBlockReference", "DataSetReference", "ApplicationId",
             "VlanId", "VlanPriority", "ConfigurationRevision", "SynchronizationStatus",
             "GrandMasterClockId", "SampleMode", "SampleRate"]


# ---------------------------------------------------------------- .ohc (MFC archive)
def cstr(s):
    """MFC unicode CString"""
    n = len(s)
    head = b"\xff\xfe\xff" + (bytes([n]) if n < 255 else b"\xff" + struct.pack("<H", n))
    return head + s.encode("utf-16le")


def astr(s):
    """ANSI string with a 32-bit length (values of the SV tables)"""
    e = s.encode("cp1252", "replace")
    return struct.pack("<I", len(e)) + e


def iter_cstr(b, a, z):
    i = a
    while True:
        i = b.find(b"\xff\xfe\xff", i, z)
        if i < 0:
            return
        n = b[i + 3]
        j = i + 4
        if n == 0xFF:
            n = struct.unpack("<H", b[j:j + 2])[0]
            j += 2
        yield i, j + 2 * n, b[j:j + 2 * n].decode("utf-16le", "replace")
        i = j + 2 * n


def section(b, name, nxt):
    a = b.find(name.encode("utf-16le"))
    if a < 0:
        raise ValueError("section %s not found in the master" % name)
    z = b.find(nxt.encode("utf-16le"), a)
    if z < 0:
        raise ValueError("end of section %s not found" % name)
    return a, z


def terminals(b, name, nxt):
    """[(start, end, name)] of the terminal names of a section (a name followed by 'Name')"""
    a, z = section(b, name, nxt)
    L = list(iter_cstr(b, a, z))
    return [L[k] for k in range(len(L) - 1) if L[k + 1][2] == "Name" and L[k][2] not in ("Name", "")]


def name_flag_pos(b, name_end):
    """v0.0.4: byte 'the name was changed by the user' of a channel record:
    <name> 'Name' '' 'Number' <00000000 FLAG ...>; Test Universe sets it to 1 when a channel is
    renamed (checked with a renamed 'Bin. in 5'). Returns its position or None."""
    for (s, e, t) in iter_cstr(b, name_end, min(len(b), name_end + 200)):
        if t == "Number":
            return e + 4 if b[e:e + 4] == b"\x00\x00\x00\x00" and b[e + 4] in (0, 1) else None
    return None


def mu_records(b):
    """3 SV streams: [{field: (start, end)}] - value of a field = between its marker and its name"""
    a, z = section(b, "Sampled_Values_Merging_Unit", "Sampled_Values_Data_Attributes")
    names = list(iter_cstr(b, a, z))
    recs, cur, k = [], None, 0
    for (s, e, nm) in names:
        if nm not in MU_FIELDS:
            continue
        idx = MU_FIELDS.index(nm)
        marker = b"\x02\x00\x00\x00" + struct.pack("<I", idx)
        m = b.rfind(marker, max(a, s - 400), s)
        if m < 0:
            raise ValueError("field %s of an SV stream not understood" % nm)
        if idx == 0:
            cur = {}
            recs.append(cur)
        cur[nm] = (m + 8, s)
    if len(recs) != 3 or any(len(r) != len(MU_FIELDS) for r in recs):
        raise ValueError("the master must have 3 SV streams (found %d)" % len(recs))
    return recs


def attr_keys(b):
    """24 data attribute rows: [(start, end, key)] - the key (channel) before 'DisplayName'"""
    a, z = section(b, "Sampled_Values_Data_Attributes", "OMICRON_HWC_CMC_AS_MASTER_CLOCK")
    out = []
    for (s, e, nm) in iter_cstr(b, a, z):
        if nm != "DisplayName":
            continue
        m = b.rfind(b"\x02\x00\x00\x00\x00\x00\x00\x00", max(a, s - 300), s)
        if m < 0:
            raise ValueError("data attribute row not understood")
        n = struct.unpack("<I", b[m + 8:m + 12])[0]
        if m + 12 + n != s:
            raise ValueError("data attribute row: unexpected layout")
        out.append((m + 8, s, b[m + 12:s].decode("cp1252", "replace")))
    return out


def _u32(b, p):
    return struct.unpack("<I", b[p:p + 4])[0]


def _skip_tag(b, p):
    """MFC class tag: new class (FFFF schema len name) or a 2-byte reference"""
    if b[p:p + 2] == b"\xff\xff":
        n = struct.unpack("<H", b[p + 4:p + 6])[0]
        return p + 6 + n
    return p + 2


def _astr_at(b, p):
    n = _u32(b, p)
    return p + 4 + n


def sv_controlblocks(b):
    """The REAL SV configuration (dialog 'Configure Sampled Values'): object SampledValues
    with 3 ControlBlocks. Each: IED, SCL path, header, 8 members (dataset attribute ->
    CMC channel), block 1 = values of the SCL, block 2 = values in use (shown in the dialog).
    Returns per stream: {'mem': [(name_start, name_end, rest_start)], 'b2': {field: (start, end)}}"""
    a = b.find(b"\x0d\x00SampledValues")
    if a < 0:
        raise ValueError("SV configuration (SampledValues) not found in the master")
    p = a + 2 + 13
    cnt = _u32(b, p)
    p += 4
    out = []
    for _ in range(cnt):
        p = _skip_tag(b, p)
        p = _astr_at(b, p)                     # IED of the SCL
        p = _astr_at(b, p)                     # SCL path
        p += 8 + 1 + 4 + 4 + 1                 # time, flags, sample rate ...
        n = _u32(b, p)
        p += 4
        mem = []
        for _ in range(n):
            p = _skip_tag(b, p)
            p = _astr_at(b, p)                 # dataset attribute (MU01.TCTR1.AmpSv.instMag.i)
            p += 4                             # 1 = current, 2 = voltage
            s0 = p
            p = _astr_at(b, p)                 # CMC channel name
            mem.append((s0, p, p))             # rest: kind, stream, channel index, 0x0d, 0, 0
            p += 24
        blocks = []
        for _ in range(2):
            f = {}
            f["svid"] = (p, _astr_at(b, p)); p = f["svid"][1]
            f["mac"] = (p, _astr_at(b, p)); p = f["mac"][1]
            f["cb"] = (p, _astr_at(b, p)); p = f["cb"][1]
            f["ds"] = (p, _astr_at(b, p)); p = f["ds"][1]
            f["nums"] = (p, p + 9); p += 9     # APPID u16, VLAN u16, prio u8, ConfRev u32
            p += 18
            blocks.append(f)
        out.append({"mem": mem, "b2": blocks[1]})
    if len(out) != 3 or any(len(o["mem"]) != 8 for o in out):
        raise ValueError("the master must have 3 SV streams with 8 attributes each")
    if b[p:p + 4] != b"\x03\x00\x00\x00":
        raise ValueError("SV configuration of the master not understood")
    return out


def apply_patches(b, reps):
    reps = sorted(reps)
    out, p = bytearray(), 0
    for s, e, nb in reps:
        if s < p:
            raise ValueError("overlapping changes")
        out += b[p:s]
        out += nb
        p = e
    out += b[p:]
    return bytes(out)


class Master:
    def __init__(self, path):
        self.path = path
        self.b = open(path, "rb").read()
        self.bi = terminals(self.b, "OMICRON_BinInpTerminal", "OMICRON_BinInpAssign")
        self.bo = terminals(self.b, "OMICRON_BinOutTerminal", "OMICRON_BinOutAssign")
        self.ao = terminals(self.b, "OMICRON_AnaOutTerminal", "OMICRON_AnaOutAssign")
        self.mu = mu_records(self.b)
        self.attr = attr_keys(self.b)
        self.cbs = sv_controlblocks(self.b)
        if len(self.ao) != 24:
            raise ValueError("the master must have 6 SV groups x 4 channels (found %d channels)" % len(self.ao))
        if len(self.attr) not in (0, 24):
            raise ValueError("the master has %d data attribute rows (0 or 24 expected)" % len(self.attr))


# ---------------------------------------------------------------- Excel
def T(v):
    return "" if v is None else str(v).strip()


def open_wb(path):
    """read-only; a file open in Excel is read from a temporary copy"""
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True, keep_vba=False), None
    except PermissionError:
        tmp = os.path.join(tempfile.gettempdir(), "omicron_cfg_" + os.path.basename(path))
        shutil.copy2(path, tmp)
        return openpyxl.load_workbook(tmp, read_only=True, data_only=True, keep_vba=False), tmp


def read_sv(gpb):
    """General_PB SV: SV ID -> dict (ied, mac, app, vlan, prio, ds)"""
    wb, tmp = open_wb(gpb)
    try:
        ws = wb["SV"]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
        if tmp:
            os.remove(tmp)
    head = [T(x).upper() for x in rows[0]]

    def col(*names):
        for n in names:
            if n.upper() in head:
                return head.index(n.upper())
        raise ValueError("General_PB SV: column %s not found" % names[0])
    cI, cS, cV, cP, cM, cA, cD = (col("IEDName"), col("SV ID"), col("VLAN ID DEC"), col("VLAN p"),
                                  col("DEST MAC ADDR"), col("APP ID DEC"), col("Dataset"))
    sv, ied = {}, ""
    for r in rows[1:]:
        r = list(r) + [None] * 20
        if T(r[cI]):
            ied = T(r[cI])
        sid = T(r[cS])
        if not sid:
            continue
        sv[sid.upper()] = dict(svid=sid, ied=ied, mac=T(r[cM]), app=r[cA], vlan=r[cV], prio=r[cP], ds=T(r[cD]))
    return sv


def read_device(path):
    """bi / bo: label -> Omicron ID; tests: [(test name, {label: (id, svid)})]"""
    wb, tmp = open_wb(path)
    try:
        dev = dict(bi={}, bo={}, tests=[])
        names = wb.sheetnames
        for sh, key in ((BI_SHEET, "bi"), (BO_SHEET, "bo")):
            if sh in names:
                for r in wb[sh].iter_rows(min_row=2, values_only=True):
                    r = list(r) + [None] * 8
                    if T(r[0]):
                        dev[key][T(r[0])] = T(r[3])
        for sh in names:
            if sh.startswith(AI_PREFIX):
                m = {}
                for r in wb[sh].iter_rows(min_row=2, values_only=True):
                    r = list(r) + [None] * 8
                    if T(r[0]):
                        m[T(r[0])] = (T(r[3]), T(r[5]))
                dev["tests"].append((sh[len(AI_PREFIX):].strip(), m))
        return dev
    finally:
        wb.close()
        if tmp:
            os.remove(tmp)


# ---------------------------------------------------------------- building one file
def ai_labels():
    """sheet labels in the order of the 24 master channels: V1, V2, V3, I1, I2, I3"""
    def grp(p, n):
        return [p + " L1", p + " L2", p + " L3", p + " E"] if n == 1 else \
               ["%s(%d)-1" % (p, n), "%s(%d)-2" % (p, n), "%s(%d)-3" % (p, n), "%s(%d)-E" % (p, n)]
    return [l for p in ("V", "I") for n in (1, 2, 3) for l in grp(p, n)]


ID_OK = re.compile(r"^[A-Za-z][A-Za-z0-9_]*$")


def build(master, dev, test, sv, notes, ied):
    b = master.b
    reps = []

    flags = set()

    def set_flag(name_end):
        fp = name_flag_pos(b, name_end)
        if fp is None:
            notes.append("%s: 'name changed' flag of a channel not found in the master" % ied)
        elif b[fp] != 1 and fp not in flags:
            flags.add(fp)
            reps.append((fp, fp + 1, b"\x01"))

    def rename(term, values, what):
        n = 0
        for (s, e, old) in term:
            new = values.get(old, "")
            if new:
                if not ID_OK.match(new):
                    notes.append("%s: %s '%s' - not only A-Z a-z 0-9 _" % (ied, what, new))
                if new != old:
                    reps.append((s, e, cstr(new)))
                    set_flag(e)
                    n += 1
        return n

    nbi = rename(master.bi, dev["bi"], "binary input")
    nbo = rename(master.bo, dev["bo"], "binary output")
    known = set(t[2] for t in master.bi) | set(t[2] for t in master.bo)
    for lab, v in list(dev["bi"].items()) + list(dev["bo"].items()):
        if v and lab not in known:
            notes.append("%s: '%s' (%s) - no such channel in the master" % (ied, lab, v))

    # analog outputs + streams
    names = [t[2] for t in master.ao]
    nai, used = 0, []
    if test:
        tname, rows = test
        labels = ai_labels()
        stream_sv = [None, None, None]
        for k, (s, e, old) in enumerate(master.ao):
            oid, sid = rows.get(labels[k], ("", ""))
            g = k % 12 // 4                                        # stream 0..2
            if sid:
                if stream_sv[g] and stream_sv[g].upper() != sid.upper():
                    notes.append("%s / %s: group %d has two SV IDs (%s, %s) - %s used"
                                 % (ied, tname, g + 1, stream_sv[g], sid, stream_sv[g]))
                elif not stream_sv[g]:
                    stream_sv[g] = sid
            if oid:
                if not ID_OK.match(oid):
                    notes.append("%s / %s: channel '%s' - not only A-Z a-z 0-9 _" % (ied, tname, oid))
                names[k] = oid
                if oid != old:
                    reps.append((s, e, cstr(oid)))
                    set_flag(e)
                    nai += 1
        for g in range(3):
            rec = master.mu[g]
            sid = stream_sv[g]
            d = sv.get(sid.upper()) if sid else None
            if sid and not d:
                notes.append("%s / %s: SV ID %s not on General_PB SV - stream %d left as in the master" % (ied, tname, sid, g + 1))
            if not d:
                continue                                           # unused stream: as in the master
            ld = d["ds"].split("/")[0] if "/" in d["ds"] else ""
            ldno = re.sub(r"\D", "", ld)
            cb = "%s%s/LLN0$MS$MSVCB%02d" % (d["ied"], ld, int(ldno)) if ldno else ""
            # v0.0.4: the table 'Merging Unit' (report) stays as in the master (as done by hand)
            # the real configuration (what the dialog shows): block 2 + the 8 attributes
            cbk = master.cbs[g]
            mac = bytes(int(x, 16) for x in re.split(r"[-:]", d["mac"]) if x) if d["mac"] else b""
            b2 = cbk["b2"]
            reps.append((b2["svid"][0], b2["svid"][1], astr(d["svid"])))
            reps.append((b2["mac"][0], b2["mac"][1], struct.pack("<I", len(mac)) + mac))
            # v0.0.4: Control block reference / Dataset reference stay as in the master (as done by hand)
            reps.append((b2["nums"][0], b2["nums"][1],
                         struct.pack("<HHBI", int(d["app"] or 0) & 0xFFFF, int(d["vlan"] or 0) & 0xFFFF,
                                     int(d["prio"] or 0) & 0xFF, CONF_REV)))
            for j, (s0, e0, r0) in enumerate(cbk["mem"]):
                k = (12 + g * 4 + j) if j < 4 else (g * 4 + j - 4)      # I channels first, then V
                kind = 1 if j < 4 else 2
                reps.append((s0, e0, astr(names[k])))
                reps.append((r0, r0 + 12, struct.pack("<III", kind, g, j % 4)))
            used.append(d["svid"])
            if not d["mac"] or not d["app"] or d["vlan"] in (None, ""):
                notes.append("%s / %s: SV ID %s - MAC / APP ID / VLAN empty on General_PB" % (ied, tname, sid))

    # v0.0.4: the table 'Data Attributes' (report rows, GUIDs) stays as in the master (as done by hand)
    if False:                                          # kept for reference (v0.0.3)
        for r, (s, e, key) in enumerate(master.attr):
            g, j = r // 8, r % 8
            k = (12 + g * 4 + j) if j < 4 else (g * 4 + j - 4)    # index in the 24 channels
            new = names[k]
            if new != master.ao[k][2] and new != key:
                reps.append((s, e, astr(new)))
    return apply_patches(b, reps), nbi, nbo, nai, used


# ---------------------------------------------------------------- project
class UserError(Exception):
    """a problem the user must fix (message in English)"""


def project_paths(prj):
    if not prj or not os.path.isdir(prj):
        raise UserError("Project folder not found: %s" % prj)
    gpb = os.path.join(prj, "General_PB.xlsm")
    if not os.path.isfile(gpb):
        c = [p for p in glob.glob(os.path.join(prj, "*.xlsm")) if "general_pb" in os.path.basename(p).lower()]
        if not c:
            raise UserError("General_PB.xlsm not found in %s" % prj)
        gpb = c[0]
    doc = os.path.join(prj, "DOC")
    if not os.path.isdir(doc):
        raise UserError("Folder DOC not found in %s" % prj)
    return gpb, doc, os.path.join(prj, "Library"), os.path.join(prj, OUT_DIR)


def list_devices(prj):
    """IED names = device files DOC\\<IED>.xlsm"""
    gpb, doc, lib, out = project_paths(prj)
    return [os.path.splitext(os.path.basename(p))[0] for p in sorted(glob.glob(os.path.join(doc, "*.xlsm")))
            if not os.path.basename(p).startswith("~$")]


def list_masters(prj):
    return sorted(os.path.basename(p) for p in glob.glob(os.path.join(prj, "Library", "*.ohc")))


def run(prj, devices, master=None, log=print):
    """builds the .ohc files of the given devices; returns (files written, report lines)"""
    gpb, doc, lib, out = project_paths(prj)
    masters = list_masters(prj)
    if master:
        masters = [m for m in masters if m.lower() == master.lower()]
    if not masters:
        raise UserError("No master .ohc file in %s" % lib)
    mpath = os.path.join(lib, masters[0])
    master_obj = Master(mpath)
    log("Master: %s" % masters[0])
    log("Reading General_PB ...")
    sv = read_sv(gpb)
    os.makedirs(out, exist_ok=True)
    rep, notes, made = [], [], 0
    for ied in devices:
        path = os.path.join(doc, ied + ".xlsm")
        log("  " + ied)
        try:
            dev = read_device(path)
        except Exception as ex:
            notes.append("%s: not read - %s" % (ied, ex))
            continue
        has_bin = any(dev["bi"].values()) or any(dev["bo"].values())
        tests = [t for t in dev["tests"] if any(v[0] or v[1] for v in t[1].values())]
        if not has_bin and not tests:
            rep.append("%-14s no Omicron data - skipped" % ied)
            continue
        jobs = [("%s_Omicron_AI_%s.ohc" % (ied, re.sub(r"[^\w\-]+", "_", t[0])), t) for t in tests] or \
               [("%s_Omicron.ohc" % ied, None)]
        folder = os.path.join(out, ied)
        os.makedirs(folder, exist_ok=True)
        for fn, t in jobs:
            try:
                data, nbi, nbo, nai, used = build(master_obj, dev, t, sv, notes, ied)
                with open(os.path.join(folder, fn), "wb") as f:
                    f.write(data)
                made += 1
                rep.append("%-44s BI %3d  BO %3d  AI %2d  SV %s" % (ied + "\\" + fn, nbi, nbo, nai, ", ".join(used) or "-"))
            except Exception as ex:
                notes.append("%s: %s - %s" % (ied, fn, ex))
    lines = ["omicron_configuration  %s" % datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
             "Project : %s" % prj, "Master  : %s" % os.path.basename(mpath), "Output  : %s" % out, ""] + rep + \
            (["", "NOTES:"] + ["  " + n for n in notes] if notes else []) + ["", "%d file(s) written." % made]
    with open(os.path.join(out, REPORT), "w", encoding="utf-8-sig") as f:
        f.write("\n".join(lines) + "\n")
    return made, lines


# ---------------------------------------------------------------- command line
def ask_folder():
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        p = filedialog.askdirectory(title="Project folder (General_PB.xlsm, DOC, Library)")
        root.destroy()
        return p
    except Exception:
        return input("Project folder: ").strip().strip('"')


def choose(items, title, allow_all=True):
    print("\n" + title)
    for i, it in enumerate(items, 1):
        print("  %2d  %s" % (i, it))
    while True:
        a = input(("Number(s) 1,3,5-7 / name" + (" / Enter = all" if allow_all else "")) + ": ").strip()
        if (not a or a.lower() in ("all", "*")) and allow_all:
            return list(items)
        sel = []
        try:
            for part in a.split(","):
                part = part.strip()
                if "-" in part and part.replace("-", "").isdigit():
                    x, y = map(int, part.split("-"))
                    sel += items[x - 1:y]
                elif part.isdigit():
                    sel.append(items[int(part) - 1])
                elif part:
                    sel += [it for it in items if part.lower() in it.lower()]
        except (ValueError, IndexError):
            sel = []
        if sel:
            return list(dict.fromkeys(sel))
        print("  nothing chosen - again")


def main():
    ap = argparse.ArgumentParser(description="OMICRON .ohc from the General_PB project")
    ap.add_argument("project", nargs="?")
    ap.add_argument("--device", action="append", help="IED name (repeatable)")
    ap.add_argument("--all", action="store_true", help="all devices, no question")
    ap.add_argument("--master", help="master .ohc in Library (file name)")
    a = ap.parse_args()
    code = 0
    try:
        prj = a.project or ask_folder()
        ieds = list_devices(prj)
        if not ieds:
            raise UserError("No device files in DOC")
        if a.device:
            sel = [i for i in ieds if i.lower() in [d.lower() for d in a.device]]
        else:
            sel = ieds if a.all else choose(ieds, "Devices in DOC:")
        master = a.master
        if not master and len(list_masters(prj)) > 1:
            master = choose(list_masters(prj), "Master file:", False)[0]
        made, lines = run(prj, sel, master)
        print("\n" + "\n".join(lines))
    except UserError as ex:
        print("\n" + str(ex))
        code = 1
    if not a.project:
        input("Enter - close")
    return code


if __name__ == "__main__":
    sys.exit(main())
