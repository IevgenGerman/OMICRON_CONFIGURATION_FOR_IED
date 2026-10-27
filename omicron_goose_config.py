# -*- coding: utf-8 -*-
"""
omicron_goose_config.py - OMICRON Test Universe GOOSE configuration (.xml) from the
General_PB project (author: German Ievgen).

  python omicron_goose_config.py                      -> asks for the project folder
  python omicron_goose_config.py "D:\\Projects\\Bucha"  -> project folder given
  options: --device W1G_D60 (repeatable) | --all | --master default_goose.xml

Project folder:
  General_PB.xlsm           sheet GEN (IEDName + GoCB name -> DEST MAC, used for the check)
  Library\\default_goose.xml  master (empty GOOSE configuration saved by Test Universe)
  CID\\*.cid (*.icd *.scd)    the SCL files of ALL devices used: the device itself and
                              every publisher it subscribes to
  DOC\\<IED>.xlsm             sheets Omicron Binary Inputs / Omicron Binary Outputs
Result (folder Omicron next to DOC, one sub-folder per IED):
  Omicron\\<IED>\\<IED>_GOOSE.xml  +  Omicron\\_goose_report.txt
Also used by the GUI ied2omicron.py (functions list_devices / run / update_files).

UPDATE of an existing GOOSE configuration from new CID files (update_files,
  CLI: --update A_GOOSE.xml [B_GOOSE.xml ...] --cid <folder with the new CID files>):
  a block is found in the CIDs by its GoCB reference (IED + LD + GCB name), else by MAC
    found        -> MAC, APPID, GoID, references, VLAN, ConfRev, Initial / Final and the data
                    set (items, order) from the CID; kept from the XML: Enable, SimulateFlag,
                    Multiplier and every CMC mapping of an item that is still there (same type)
    CID of the device present, GCB not -> the block is removed
    no CID of the device               -> the block is left unchanged
  written only when something changed: <name>_updated_<date>_<time>.xml next to the
  original (the original is not changed) + _goose_update_report_<date>_<time>.txt

What is built:
  GOOSESubscriptions (the CMC receives - Omicron Binary Inputs):
      every GCB of the device that appears on the sheet (column G Dest MAC);
      MAC, APPID, GoID, GoCB reference and the WHOLE data set from the CID;
      the row "Bin. in n" / "vBin. in g-k" maps its ELEMENT (column B) to CMC input n / 24+(g-1)*24+k
  GOOSESimulations (the CMC publishes - Omicron Binary Outputs):
      every GCB of another device that appears on the sheet (columns E IED Name, F GCB Name);
      MAC, APPID, GoID, GoCB reference, VLAN ID / priority, ConfRev, MinTime / MaxTime
      (repetition initial / final) and the WHOLE data set from the CID;
      the row "Bin. out n" / "vBin. out g-k" maps its ELEMENT to CMC output n / 24+(g-1)*24+k
  A GCB is found in the CID files by its MAC address (GEN of General_PB gives the MAC of
  IED Name + GCB Name), so the IED name inside the CID may differ from the project name.
  v0.0.4 - every IED ONLY FROM ITS OWN FILE: a CID / ICD / IID gives its first IED only
  (the other IEDs inside it are copies of publishers kept for the subscriptions - no VLAN,
  default repetition times - and are never used); an SCD gives all its IEDs.
  All GOOSE parameters (MAC, APPID, GoID, GoCB reference, VLAN ID / priority, ConfRev,
  MinTime / MaxTime, data set) come from that own file; General_PB GEN only gives the MAC to
  find it (a difference CID / General_PB -> note in the report).
  A device that needs a GCB without the own CID file of its IED (or without a MAC in GEN)
  is NOT created: message + report "NOT CREATED - reason"; the other devices are made.
Needs: Python 3, openpyxl (pip install openpyxl).
"""
import os, re, sys, glob, shutil, tempfile, datetime, argparse, warnings
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape, quoteattr

warnings.filterwarnings("ignore")
try:
    import openpyxl
except ImportError:
    print("openpyxl is missing: pip install openpyxl")
    sys.exit(1)

BI_SHEET = "Omicron Binary Inputs"
BO_SHEET = "Omicron Binary Outputs"
OUT_DIR = "Omicron"                 # next to DOC; one sub-folder per IED
REPORT = "_goose_report.txt"
CID_SUB = "CID"
MASTER_NAME = "default_goose.xml"
VIRT_OFFSET = 24            # vBin. in / out g-k -> 24 + (g-1)*24 + k
SUB_FLAG = "-1"             # SimulateFlag of a subscription (as Test Universe writes it)
SIM_FLAG = "2"              # SimulateFlag of a simulation
REP_MULTIPLIER = "2"

# SCL basic type -> (Omicron item type, extra element)
ITEM_TYPES = {
    "BOOLEAN": ("GOOSEBoolean", None),
    "Quality": ("GOOSEBitString", "<Length>13</Length>"),
}


def T(v):
    return "" if v is None else str(v).strip()


def mac_key(m):
    return "-".join(x.upper().zfill(2) for x in re.split(r"[-:]", T(m)) if x)


# ---------------------------------------------------------------- SCL (CID)
class Scl:
    """GCBs of one SCL file, indexed by MAC address"""

    def __init__(self, path):
        self.path = path
        r = ET.parse(path).getroot()
        self.ns = r.tag.split("}")[0][1:] if r.tag.startswith("{") else ""
        q = self.q
        dtt = r.find(q("DataTypeTemplates"))
        self.lnt = {e.get("id"): e for e in dtt.findall(q("LNodeType"))} if dtt is not None else {}
        self.dot = {e.get("id"): e for e in dtt.findall(q("DOType"))} if dtt is not None else {}
        self.dat = {e.get("id"): e for e in dtt.findall(q("DAType"))} if dtt is not None else {}
        ied_list = [i.get("name") for i in r.findall(q("IED"))]
        self.ieds = {i.get("name"): i for i in r.findall(q("IED"))}
        # v0.0.4: every IED only from ITS OWN file. A CID / ICD / IID describes one IED (the first);
        # the other IEDs inside it are copies of publishers kept for the subscriptions (no VLAN,
        # default repetition times) and are never used. An SCD owns all its IEDs.
        if os.path.splitext(path)[1].lower() == ".scd":
            self.own = set(ied_list)
        else:
            self.own = set(ied_list[:1])
        self.gcbs = {}
        self.copies = {}                     # MAC -> IED name of a copy (for the messages)
        for cap in r.iter(q("ConnectedAP")):
            ied = cap.get("iedName")
            for g in cap.findall(q("GSE")):
                adr = {p.get("type"): T(p.text) for p in g.iter(q("P"))}
                mn, mx = g.find(q("MinTime")), g.find(q("MaxTime"))
                mac = mac_key(adr.get("MAC-Address", ""))
                if mac and ied not in self.own:
                    self.copies.setdefault(mac, ied)
                elif mac:
                    self.gcbs[mac] = dict(scl=self, ied=ied, ld=g.get("ldInst"), cb=g.get("cbName"),
                                          appid=adr.get("APPID", "0"), vlan=adr.get("VLAN-ID", "0"),
                                          prio=adr.get("VLAN-PRIORITY", "0"),
                                          tmin=T(mn.text) if mn is not None else "",
                                          tmax=T(mx.text) if mx is not None else "", mac=mac)

    def q(self, t):
        return "{%s}%s" % (self.ns, t) if self.ns else t

    def control(self, g):
        """GSEControl + data set of a GCB"""
        ied = self.ieds.get(g["ied"])
        if ied is None:
            raise ValueError("IED %s not in %s" % (g["ied"], os.path.basename(self.path)))
        for ld in ied.iter(self.q("LDevice")):
            if ld.get("inst") != g["ld"]:
                continue
            ln0 = ld.find(self.q("LN0"))
            for gc in ln0.findall(self.q("GSEControl")):
                if gc.get("name") == g["cb"]:
                    ds = [d for d in ln0.findall(self.q("DataSet")) if d.get("name") == gc.get("datSet")]
                    return gc, (ds[0] if ds else None)
        raise ValueError("GSEControl %s/%s not in %s" % (g["ld"], g["cb"], os.path.basename(self.path)))

    def _ln_type(self, ied, ldi, lnn):
        for ld in ied.iter(self.q("LDevice")):
            if ld.get("inst") == ldi:
                for ln in ld:
                    if ln.tag in (self.q("LN0"), self.q("LN")) and \
                            (ln.get("prefix") or "") + ln.get("lnClass") + (ln.get("inst") or "") == lnn:
                        return ln.get("lnType")
        raise ValueError("LN %s/%s not found" % (ldi, lnn))

    def _leaves(self, typ, path, fc, kind):
        out = []
        q = self.q
        if kind == "DO":
            for ch in self.dot[typ]:
                if ch.tag == q("DA"):
                    if fc and ch.get("fc") != fc:
                        continue
                    p = path + "." + ch.get("name")
                    if ch.get("bType") == "Struct":
                        out += self._leaves(ch.get("type"), p, fc, "DA")
                    else:
                        out.append((p, ch.get("bType")))
                elif ch.tag == q("SDO"):
                    out += self._leaves(ch.get("type"), path + "." + ch.get("name"), fc, "DO")
        else:
            for ch in self.dat[typ]:
                if ch.tag != q("BDA"):
                    continue
                p = path + "." + ch.get("name")
                if ch.get("bType") == "Struct":
                    out += self._leaves(ch.get("type"), p, fc, "DA")
                else:
                    out.append((p, ch.get("bType")))
        return out

    def items(self, g, ds):
        """[(item name, SCL basic type)] of the data set, in order"""
        ied = self.ieds[g["ied"]]
        out = []
        for f in ds.findall(self.q("FCDA")):
            lnn = (f.get("prefix") or "") + f.get("lnClass") + (f.get("lnInst") or "")
            ldi = f.get("ldInst")
            lt = self._ln_type(ied, ldi, lnn)
            do = f.get("doName").split(".")
            dtyp = [x.get("type") for x in self.lnt[lt] if x.get("name") == do[0]][0]
            for sd in do[1:]:
                dtyp = [x.get("type") for x in self.dot[dtyp] if x.get("name") == sd][0]
            base = "%s%s/%s.%s" % (g["ied"], ldi, lnn, f.get("doName"))
            da = f.get("daName")
            if da:
                parts = da.split(".")
                el = [x for x in self.dot[dtyp] if x.get("name") == parts[0]][0]
                for pp in parts[1:]:
                    el = [x for x in self.dat[el.get("type")] if x.get("name") == pp][0]
                if el.get("bType") == "Struct":
                    out += self._leaves(el.get("type"), base + "." + da, None, "DA")
                else:
                    out.append((base + "." + da, el.get("bType")))
            else:
                out += self._leaves(dtyp, base, f.get("fc"), "DO")
        return out


def load_scl(folder, copies=None):
    """GCBs of the OWN IEDs of all SCL files, by MAC. copies (dict, optional) gets
    MAC -> [(IED, file)] of the copies inside other files (only for the messages)."""
    gcbs, errs = {}, []
    files = [p for p in glob.glob(os.path.join(folder, "*")) if os.path.splitext(p)[1].lower() in (".cid", ".icd", ".scd", ".iid")]
    for p in sorted(files):
        try:
            s = Scl(p)
            for mac, g in s.gcbs.items():
                if mac in gcbs and gcbs[mac]["ied"] != g["ied"]:
                    errs.append("MAC %s: own GCB of %s (%s) and of %s (%s) - the first is used"
                                % (mac, gcbs[mac]["ied"], os.path.basename(gcbs[mac]["scl"].path), g["ied"],
                                   os.path.basename(p)))
                gcbs.setdefault(mac, g)
            if copies is not None:
                for mac, ied in s.copies.items():
                    copies.setdefault(mac, []).append((ied, os.path.basename(p)))
        except Exception as ex:
            errs.append("%s: not read - %s" % (os.path.basename(p), ex))
    return gcbs, files, errs


# ---------------------------------------------------------------- Excel
def open_wb(path):
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True, keep_vba=False), None
    except PermissionError:
        tmp = os.path.join(tempfile.gettempdir(), "omicron_goose_" + os.path.basename(path))
        shutil.copy2(path, tmp)
        return openpyxl.load_workbook(tmp, read_only=True, data_only=True, keep_vba=False), tmp


def read_gen(gpb):
    """GEN: (IEDName, GoCB name) -> dict(mac, goid, app, vlan)"""
    wb, tmp = open_wb(gpb)
    try:
        rows = list(wb["GEN"].iter_rows(values_only=True))
    finally:
        wb.close()
        if tmp:
            os.remove(tmp)
    head = [T(x).upper() for x in rows[0]]

    def col(n):
        if n.upper() not in head:
            raise ValueError("General_PB GEN: column %s not found" % n)
        return head.index(n.upper())
    cI, cG, cO, cM, cA, cV = col("IEDName"), col("GoCB name"), col("GOID"), col("DEST MAC ADDR"), col("APP ID DEC"), col("VLAN ID")
    gen, ied = {}, ""
    for r in rows[1:]:
        r = list(r) + [None] * 30
        if T(r[cI]):
            ied = T(r[cI])
        if T(r[cG]) and ied:
            gen[(ied.upper(), T(r[cG]).upper())] = dict(mac=mac_key(r[cM]), goid=T(r[cO]), app=T(r[cA]), vlan=T(r[cV]))
    return gen


def read_device(path):
    """rows of the two Omicron sheets: (label, element, omicron id, ied, gcb, mac)"""
    wb, tmp = open_wb(path)
    try:
        dev = dict(bi=[], bo=[])
        for sh, key in ((BI_SHEET, "bi"), (BO_SHEET, "bo")):
            if sh in wb.sheetnames:
                for r in wb[sh].iter_rows(min_row=2, values_only=True):
                    r = list(r) + [None] * 10
                    if T(r[0]) and T(r[1]):
                        dev[key].append((T(r[0]), T(r[1]), T(r[3]), T(r[4]), T(r[5]), mac_key(r[6]) if key == "bi" else ""))
        return dev
    finally:
        wb.close()
        if tmp:
            os.remove(tmp)


# ---------------------------------------------------------------- mapping helpers
LABEL = re.compile(r"^(v?)Bin\.\s*(?:in|out)\s+(\d+)(?:-(\d+))?$", re.I)


def io_number(label):
    m = LABEL.match(label.strip())
    if not m:
        return None
    if m.group(1):
        if not m.group(3):
            return None
        return VIRT_OFFSET + (int(m.group(2)) - 1) * 24 + int(m.group(3))
    return int(m.group(2))


def element_key(el):
    """'Gen.FlxLgcGAPC1.ST.Ind053.stVal' / 'CTRL/BININGGIO1.ST.Ind25.stVal' -> 'GEN/FLXLGCGAPC1.IND053.STVAL'"""
    p = el.replace("/", ".", 1).split(".")
    if len(p) < 3:
        return el.upper()
    rest = p[3:] if re.fullmatch(r"[A-Z]{2}", p[2]) else p[2:]
    return (p[0] + "/" + p[1] + "." + ".".join(rest)).upper()


def item_key(name, ied):
    return name[len(ied):].upper() if name.startswith(ied) else name.upper()


# ---------------------------------------------------------------- XML
def item_xml(name, btype, ios):
    """ios: CMC numbers (new mapping) or dicts (mapping kept from an existing XML)"""
    typ, extra = ITEM_TYPES[btype]
    L = ['        <Item xsi:type="%s" Name=%s>' % (typ, quoteattr(name))]
    if ios:
        L.append("          <CMCMappings>")
        for n in ios:
            m = n if isinstance(n, dict) else dict(Name="", IONumber=str(n), Bit="0", Value="0", Inverted="false")
            L += ['            <CMCMapping Name=%s>' % quoteattr(m.get("Name", "")),
                  "              <IONumber>%s</IONumber>" % m.get("IONumber", "0"),
                  "              <Bit>%s</Bit>" % m.get("Bit", "0"),
                  "              <Value>%s</Value>" % m.get("Value", "0"),
                  "              <Inverted>%s</Inverted>" % m.get("Inverted", "false"),
                  "            </CMCMapping>"]
        L.append("          </CMCMappings>")
    else:
        L.append("          <CMCMappings />")
    if extra:
        L.append("          " + extra)
    L.append("        </Item>")
    return L


def gcb_xml(tag, g, gc, items, maps, simulate, enable="true", simflag=None, multiplier=None):
    ied = g["ied"]
    L = ['    <%s Enable="%s">' % (tag, enable),
         "      <SimulateFlag>%s</SimulateFlag>" % (simflag if simflag is not None else (SIM_FLAG if simulate else SUB_FLAG)),
         "      <MACAddress>%s</MACAddress>" % g["mac"],
         "      <ApplicationID>%d</ApplicationID>" % int(g["appid"] or "0", 16),
         "      <GOOSEID>%s</GOOSEID>" % escape(gc.get("appID") or ""),
         "      <GOOSEControlBlockReference>%s</GOOSEControlBlockReference>" % escape("%s%s/LLN0$GO$%s" % (ied, g["ld"], g["cb"]))]
    if simulate:
        L += ["      <VLAN>",
              "        <Priority>%d</Priority>" % int(g["prio"] or "0"),
              "        <ID>%d</ID>" % int(g["vlan"] or "0", 16),
              "      </VLAN>",
              "      <ConfigurationRevision>%s</ConfigurationRevision>" % (gc.get("confRev") or "1"),
              "      <Repetition>",
              "        <Final>%s</Final>" % (g["tmax"] or "1000"),
              "        <Initial>%s</Initial>" % (g["tmin"] or "4"),
              "        <Multiplier>%s</Multiplier>" % (multiplier or REP_MULTIPLIER),
              "      </Repetition>"]
    L.append("      <DataSet>")
    for name, bt in items:
        L += item_xml(name, bt, maps.get(item_key(name, ied), []))
    L.append("      </DataSet>")
    L.append("      <DataSetName>%s</DataSetName>" % escape("%s%s/LLN0$%s" % (ied, g["ld"], gc.get("datSet"))))
    L.append("    </%s>" % tag)
    return L


def read_master(path):
    raw = open(path, "rb").read()
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        enc = "utf-16"
    elif raw.startswith(b"\xef\xbb\xbf"):
        enc = "utf-8-sig"
    else:
        enc = "utf-8"
    txt = raw.decode(enc).replace("\r\n", "\n")
    if "<GOOSEConfig" not in txt:
        raise ValueError("%s is not a GOOSE configuration" % path)
    return txt, enc


def fill_master(txt, subs, sims):
    def block(tag, body):
        return "  <%s />" % tag if not body else "  <%s>\n%s\n  </%s>" % (tag, "\n".join(body), tag)
    for tag, body in (("GOOSESubscriptions", subs), ("GOOSESimulations", sims)):
        new = block(tag, body)
        txt, n = re.subn(r"  <%s\s*/>|  <%s>.*?</%s>" % (tag, tag, tag), lambda m: new, txt, count=1, flags=re.S)
        if n == 0:
            raise ValueError("<%s> not found in the master" % tag)
    return txt


# ---------------------------------------------------------------- one device
def plan_device(ied, dev, gen, gcbs, copies=None):
    """GCBs needed by a device: [(kind, mac, label for messages, rows)], missing [(text)]"""
    plan, missing = [], []
    subs = {}
    for row in dev["bi"]:
        mac = row[5]
        if not mac and row[3] and row[4]:
            mac = gen.get((row[3].upper(), row[4].upper()), {}).get("mac", "")
        subs.setdefault(mac or "?%s/%s" % (row[3], row[4]), []).append(row)
    sims = {}
    for row in dev["bo"]:
        g = gen.get((row[3].upper(), row[4].upper()))
        mac = g["mac"] if g else ""
        if not mac:
            # project name = CID name: take the GCB by IED + cbName
            for m, x in gcbs.items():
                if x["ied"].upper() == row[3].upper() and (x["cb"] or "").upper() == row[4].upper():
                    mac = m
        sims.setdefault(mac or "?%s/%s" % (row[3], row[4]), []).append(row)
    for kind, d in (("sub", subs), ("sim", sims)):
        for mac, rows in d.items():
            what = "%s / GCB %s" % (rows[0][3] or ied, rows[0][4])
            pub = rows[0][3] or ied
            if mac.startswith("?"):
                missing.append("%s / GCB %s - no MAC address for it in General_PB GEN" % (pub, rows[0][4]))
            elif mac not in gcbs:
                cp = (copies or {}).get(mac)
                if cp:
                    missing.append("%s / GCB %s (MAC %s) - no own CID file of %s in the CID folder; it is "
                                   "only a copy inside %s (a copy has no VLAN / repetition times, it is not used)"
                                   % (pub, rows[0][4], mac, pub, ", ".join(sorted(set(f for _, f in cp)))))
                else:
                    missing.append("%s / GCB %s (MAC %s) - no CID file of %s in the CID folder"
                                   % (pub, rows[0][4], mac, pub))
            else:
                plan.append((kind, mac, what, rows))
    return plan, missing


def build_device(ied, plan, gcbs, gen, notes):
    subs, sims, stat = [], [], []
    for kind, mac, what, rows in plan:
        g = gcbs[mac]
        scl = g["scl"]
        gc, ds = scl.control(g)
        if ds is None:
            notes.append("%s: %s - GSEControl without a data set, skipped" % (ied, what))
            continue
        items = scl.items(g, ds)
        bad = sorted(set(bt for _, bt in items if bt not in ITEM_TYPES))
        if bad:
            notes.append("%s: %s - data type(s) %s not supported yet (send a sample), GCB skipped"
                         % (ied, what, ", ".join(bad)))
            continue
        keys = set(item_key(n, g["ied"]) for n, _ in items)
        maps, nmap = {}, 0
        for (label, el, oid, ri, rg, rmac) in rows:
            n = io_number(label)
            k = element_key(el)
            if n is None:
                notes.append("%s: row '%s' - not a Bin. in / out label" % (ied, label))
            elif k not in keys:
                notes.append("%s: %s - element %s (%s) not in the data set %s"
                             % (ied, what, el, label, gc.get("datSet")))
            else:
                maps.setdefault(k, []).append(n)
                nmap += 1
        # project data vs CID
        gp = None
        for (gi, gg), v in gen.items():
            if v["mac"] == mac:
                gp = v
        if gp:
            diff = []
            if gp["goid"] and gp["goid"] != (gc.get("appID") or ""):
                diff.append("GoID %s / %s" % (gp["goid"], gc.get("appID")))
            if gp["app"] and str(int(float(gp["app"]))) != str(int(g["appid"] or "0", 16)):
                diff.append("APPID %s / %d" % (gp["app"], int(g["appid"] or "0", 16)))
            if gp["vlan"] and str(int(float(gp["vlan"]))) != str(int(g["vlan"] or "0", 16)):
                diff.append("VLAN %s / %d" % (gp["vlan"], int(g["vlan"] or "0", 16)))
            if diff:
                notes.append("%s: %s - General_PB / CID differ: %s (the CID is used)" % (ied, what, "; ".join(diff)))
        x = gcb_xml("GOOSESimulation" if kind == "sim" else "GOOSESubscription", g, gc, items, maps, kind == "sim")
        (sims if kind == "sim" else subs).extend(x)
        stat.append("%s %-28s %-26s %3d item(s), %3d mapped  [%s]"
                    % ("SIM" if kind == "sim" else "SUB", gc.get("appID"), g["mac"], len(items), nmap,
                       os.path.basename(scl.path)))
    return subs, sims, stat


# ---------------------------------------------------------------- project
class UserError(Exception):
    """a problem the user must fix (message in English)"""


class MissingCid(UserError):
    pass


def project_paths(prj):
    if not prj or not os.path.isdir(prj):
        raise UserError("Project folder not found: %s" % prj)
    gpb = os.path.join(prj, "General_PB.xlsm")
    if not os.path.isfile(gpb):
        c = [p for p in glob.glob(os.path.join(prj, "*.xlsm")) if "general_pb" in os.path.basename(p).lower()]
        if not c:
            raise UserError("General_PB.xlsm not found in %s" % prj)
        gpb = c[0]
    doc = os.path.join(prj, "DOC")
    if not os.path.isdir(doc):
        raise UserError("Folder DOC not found in %s" % prj)
    return gpb, doc, os.path.join(prj, CID_SUB), os.path.join(prj, OUT_DIR)


def list_devices(prj):
    """IED names = device files DOC\\<IED>.xlsm"""
    gpb, doc, cid, out = project_paths(prj)
    return [os.path.splitext(os.path.basename(p))[0] for p in sorted(glob.glob(os.path.join(doc, "*.xlsm")))
            if not os.path.basename(p).startswith("~$")]


def run(prj, devices, master=MASTER_NAME, log=print):
    """builds <IED>_GOOSE.xml of the given devices; returns (files written, report lines, problems).
    v0.0.4: every GCB only from the own CID file of its IED. A device with a GCB that has no own
    CID file (or no MAC in General_PB) is NOT created: problems = ["<IED>: NOT CREATED - reason", ...]."""
    gpb, doc, cid, out = project_paths(prj)
    mpath = os.path.join(prj, "Library", master or MASTER_NAME)
    if not os.path.isfile(mpath):
        raise UserError("Master file not found: %s" % mpath)
    if not os.path.isdir(cid):
        raise UserError("Folder with the CID files not found: %s" % cid)
    log("Reading General_PB, CID files and device files ...")
    gen = read_gen(gpb)
    copies = {}
    gcbs, cidfiles, cerr = load_scl(cid, copies)
    plans, notes, probs = {}, list(cerr), []
    for ied in devices:
        try:
            dev = read_device(os.path.join(doc, ied + ".xlsm"))
        except Exception as ex:
            notes.append("%s: not read - %s" % (ied, ex))
            continue
        if not dev["bi"] and not dev["bo"]:
            notes.append("%s: no Omicron Binary Inputs / Outputs - skipped" % ied)
            continue
        plan, miss = plan_device(ied, dev, gen, gcbs, copies)
        if miss:
            # ---- the device is NOT created: the reasons go to the message and the report
            for m in miss:
                probs.append("%s: NOT CREATED - %s" % (ied, m))
            continue
        plans[ied] = plan
    txt, enc = read_master(mpath)
    os.makedirs(out, exist_ok=True)
    rep, made = [], 0
    for ied, plan in plans.items():
        log("  " + ied)
        try:
            subs, sims, stat = build_device(ied, plan, gcbs, gen, notes)
            data = fill_master(txt, subs, sims).replace("\n", "\r\n")
            folder = os.path.join(out, ied)
            os.makedirs(folder, exist_ok=True)
            fn = "%s_GOOSE.xml" % ied
            with open(os.path.join(folder, fn), "wb") as f:
                f.write(data.encode(enc))
            made += 1
            rep.append(ied + "\\" + fn)
            rep += ["    " + s for s in stat]
        except Exception as ex:
            notes.append("%s: %s" % (ied, ex))
    lines = ["omicron_goose_config  %s" % datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
             "Project : %s" % prj, "Master  : %s" % os.path.basename(mpath),
             "CID     : %d file(s): %s" % (len(cidfiles), ", ".join(os.path.basename(p) for p in cidfiles) or "-"),
             "Output  : %s" % out, ""] + rep + \
            (["", "NOT CREATED (put the CID file of the IED into the CID folder and run again):"] +
             ["  " + p for p in probs] if probs else []) + \
            (["", "NOTES:"] + ["  " + n for n in notes] if notes else []) + \
            ["", "%d file(s) written, %d device(s) not created." % (made, len(set(p.split(":")[0] for p in probs)))]
    with open(os.path.join(out, REPORT), "w", encoding="utf-8-sig") as f:
        f.write("\n".join(lines) + "\n")
    return made, lines, probs


# ---------------------------------------------------------------- UPDATE an existing GOOSE configuration
XSI = "http://www.w3.org/2001/XMLSchema-instance"
BLOCK = re.compile(r"    <(GOOSESubscription|GOOSESimulation)\b[^>]*>.*?</\1>", re.S)
UPD_TAG = re.compile(r"_updated_\d{4}-\d{2}-\d{2}_\d{4,6}$")
OM_TYPE = {v[0]: k for k, v in ITEM_TYPES.items()}       # GOOSEBoolean -> BOOLEAN ...


def _txt(el, path, default=""):
    x = el.find(path)
    return T(x.text) if x is not None and x.text is not None else default


def parse_block(text):
    """one <GOOSESubscription> / <GOOSESimulation> of an existing file"""
    r = ET.fromstring('<r xmlns:xsi="%s">%s</r>' % (XSI, text))[0]
    b = dict(tag=r.tag, enable=r.get("Enable", "true"), simflag=_txt(r, "SimulateFlag"),
             mac=mac_key(_txt(r, "MACAddress")), app=_txt(r, "ApplicationID"), goid=_txt(r, "GOOSEID"),
             ref=_txt(r, "GOOSEControlBlockReference"), dsname=_txt(r, "DataSetName"),
             prio=_txt(r, "VLAN/Priority"), vlan=_txt(r, "VLAN/ID"), confrev=_txt(r, "ConfigurationRevision"),
             final=_txt(r, "Repetition/Final"), initial=_txt(r, "Repetition/Initial"),
             multiplier=_txt(r, "Repetition/Multiplier"), items=[])
    for it in r.findall("DataSet/Item"):
        maps = [dict(Name=m.get("Name", ""), IONumber=_txt(m, "IONumber", "0"), Bit=_txt(m, "Bit", "0"),
                     Value=_txt(m, "Value", "0"), Inverted=_txt(m, "Inverted", "false"))
                for m in it.findall("CMCMappings/CMCMapping")]
        b["items"].append((it.get("Name", ""), it.get("{%s}type" % XSI, ""), maps))
    return b


def _ref(g):
    return "%s%s/LLN0$GO$%s" % (g["ied"], g["ld"], g["cb"])


def update_files(xml_paths, cid_folder, log=print):
    """updates existing <IED>_GOOSE.xml files from new CID files; every changed file is written
    next to its original as <name>_updated_<date>_<time>.xml (the original is not changed).
    Returns (files written, report lines)."""
    if not os.path.isdir(cid_folder):
        raise UserError("Folder with the CID files not found: %s" % cid_folder)
    log("Reading CID files ...")
    gcbs, cidfiles, cerr = load_scl(cid_folder)
    if not cidfiles:
        raise UserError("No CID files in %s" % cid_folder)
    byref = {_ref(g): g for g in gcbs.values()}
    lds = {}                                            # IED name in the new CIDs -> its LD instances
    for g in gcbs.values():
        ied = g["scl"].ieds.get(g["ied"])
        if ied is not None and g["ied"] not in lds:
            lds[g["ied"]] = set(ld.get("inst") for ld in ied.iter(g["scl"].q("LDevice")))

    def device_in_cids(ref):
        head = ref.split("/LLN0$")[0]
        return any(head.startswith(i) and head[len(i):] in l for i, l in lds.items())

    stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H%M%S")
    rep, notes, made = [], list(cerr), 0
    for path in xml_paths:
        name = os.path.basename(path)
        log("  " + name)
        try:
            txt, enc = read_master(path)
        except Exception as ex:
            notes.append("%s: not read - %s" % (name, ex))
            continue
        sections = {"GOOSESubscription": [], "GOOSESimulation": []}
        lines, changed = [], False
        for m in BLOCK.finditer(txt):
            old_text = m.group(0)
            b = parse_block(old_text)
            kind = "SIM" if b["tag"] == "GOOSESimulation" else "SUB"
            label = "%s %s" % (kind, b["goid"] or b["ref"])
            g = byref.get(b["ref"]) or gcbs.get(b["mac"])
            if g is None:
                if device_in_cids(b["ref"]):
                    lines.append("    %s - GCB removed (not in the new CID)" % label)
                    changed = True
                else:
                    lines.append("    %s - no CID for this device, left unchanged" % label)
                    sections[b["tag"]].append(old_text)
                continue
            scl = g["scl"]
            try:
                gc, ds = scl.control(g)
                if ds is None:
                    raise ValueError("GSEControl without a data set")
                items = scl.items(g, ds)
            except Exception as ex:
                lines.append("    %s - CID not understood (%s), left unchanged" % (label, ex))
                sections[b["tag"]].append(old_text)
                continue
            bad = sorted(set(bt for _, bt in items if bt not in ITEM_TYPES))
            if bad:
                lines.append("    %s - data type(s) %s not supported yet, left unchanged" % (label, ", ".join(bad)))
                sections[b["tag"]].append(old_text)
                continue
            # old IED name (the CID may rename the IED): prefix of the old reference
            tail = "%s/LLN0$GO$%s" % (g["ld"], g["cb"])
            old_ied = b["ref"][:-len(tail)] if b["ref"].endswith(tail) else g["ied"]
            old = {}
            for (n, t, mp) in b["items"]:
                old[item_key(n, old_ied)] = (n, t, mp)
            new_keys = set()
            maps, msgs = {}, []
            for (n, bt) in items:
                k = item_key(n, g["ied"])
                new_keys.add(k)
                if k in old:
                    on, ot, omp = old[k]
                    if omp and OM_TYPE.get(ot) != bt:
                        msgs.append("type changed %s (%s -> %s), mapping to %s removed"
                                    % (n, ot, ITEM_TYPES[bt][0], ", ".join(x["IONumber"] for x in omp)))
                    elif omp:
                        maps[k] = omp
                else:
                    msgs.append("added %s" % n)
            for k, (on, ot, omp) in old.items():
                if k not in new_keys:
                    msgs.append("removed %s%s" % (on, (", mapping to %s lost" % ", ".join(x["IONumber"] for x in omp)) if omp else ""))
            simulate = b["tag"] == "GOOSESimulation"
            new_text = "\n".join(gcb_xml(b["tag"], g, gc, items, maps, simulate, enable=b["enable"],
                                         simflag=b["simflag"] or None, multiplier=b["multiplier"] or None))
            if re.sub(r"\s+", "", new_text) == re.sub(r"\s+", "", old_text):
                lines.append("    %s - up to date" % label)
                sections[b["tag"]].append(old_text)
                continue
            nb = parse_block(new_text)
            for f, cap in (("mac", "MAC"), ("app", "APPID"), ("goid", "GoID"), ("ref", "GoCB ref"), ("dsname", "DataSet"),
                           ("prio", "VLAN priority"), ("vlan", "VLAN ID"), ("confrev", "ConfRev"),
                           ("initial", "Initial"), ("final", "Final")):
                if b[f] != nb[f]:
                    msgs.insert(0, "%s %s -> %s" % (cap, b[f] or "-", nb[f] or "-"))
            lines.append("    %s - UPDATED [%s]" % (label, os.path.basename(scl.path)))
            lines += ["        " + x for x in (msgs or ["(layout only)"])]
            sections[b["tag"]].append(new_text)
            changed = True
        if not changed:
            rep.append("%s - nothing to update, not written" % name)
            rep += lines
            continue
        out_txt = fill_master(txt, sections["GOOSESubscription"], sections["GOOSESimulation"]).replace("\n", "\r\n")
        stem = UPD_TAG.sub("", os.path.splitext(name)[0])
        out = os.path.join(os.path.dirname(path), "%s_updated_%s.xml" % (stem, stamp))
        with open(out, "wb") as f:
            f.write(out_txt.encode(enc))
        made += 1
        rep.append("%s -> %s" % (name, os.path.basename(out)))
        rep += lines
    folder = os.path.dirname(xml_paths[0]) if xml_paths else cid_folder
    lines = ["omicron_goose_config - UPDATE  %s" % datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
             "CID     : %s (%d file(s))" % (cid_folder, len(cidfiles)), ""] + rep + \
            (["", "NOTES:"] + ["  " + n for n in notes] if notes else []) + ["", "%d file(s) written." % made]
    with open(os.path.join(folder, "_goose_update_report_%s.txt" % stamp), "w", encoding="utf-8-sig") as f:
        f.write("\n".join(lines) + "\n")
    return made, lines


# ---------------------------------------------------------------- command line
def ask_folder():
    try:
        import tkinter as tk
        from tkinter import filedialog
        root = tk.Tk()
        root.withdraw()
        p = filedialog.askdirectory(title="Project folder (General_PB.xlsm, DOC, CID, Library)")
        root.destroy()
        return p
    except Exception:
        return input("Project folder: ").strip().strip('"')


def choose(items, title):
    print("\n" + title)
    for i, it in enumerate(items, 1):
        print("  %2d  %s" % (i, it))
    while True:
        a = input("Number(s) 1,3,5-7 / name / Enter = all: ").strip()
        if not a or a.lower() in ("all", "*"):
            return list(items)
        sel = []
        try:
            for part in a.split(","):
                part = part.strip()
                if "-" in part and part.replace("-", "").isdigit():
                    x, y = map(int, part.split("-"))
                    sel += items[x - 1:y]
                elif part.isdigit():
                    sel.append(items[int(part) - 1])
                elif part:
                    sel += [it for it in items if part.lower() in it.lower()]
        except (ValueError, IndexError):
            sel = []
        if sel:
            return list(dict.fromkeys(sel))
        print("  nothing chosen - again")


def main():
    ap = argparse.ArgumentParser(description="OMICRON GOOSE configuration from the General_PB project")
    ap.add_argument("project", nargs="?")
    ap.add_argument("--device", action="append", help="IED name (repeatable)")
    ap.add_argument("--all", action="store_true", help="all devices, no question")
    ap.add_argument("--master", default=MASTER_NAME, help="master file in Library (default %s)" % MASTER_NAME)
    ap.add_argument("--update", nargs="+", metavar="XML", help="update existing GOOSE .xml file(s) from new CID files")
    ap.add_argument("--cid", help="folder with the new CID files (with --update)")
    a = ap.parse_args()
    code = 0
    if a.update:
        try:
            if not a.cid:
                raise UserError("--update needs --cid <folder with the new CID files>")
            made, lines = update_files(a.update, a.cid)
            print("\n" + "\n".join(lines))
        except UserError as ex:
            print("\n" + str(ex))
            code = 1
        return code
    try:
        prj = a.project or ask_folder()
        ieds = list_devices(prj)
        if not ieds:
            raise UserError("No device files in DOC")
        if a.device:
            sel = [i for i in ieds if i.lower() in [d.lower() for d in a.device]]
        else:
            sel = ieds if a.all else choose(ieds, "Devices in DOC:")
        made, lines, probs = run(prj, sel, a.master)
        print("\n" + "\n".join(lines))
        if probs:
            code = 2
    except MissingCid as ex:
        print("\n" + str(ex))
        code = 2
    except UserError as ex:
        print("\n" + str(ex))
        code = 1
    if not a.project:
        input("Enter - close")
    return code


if __name__ == "__main__":
    sys.exit(main())
